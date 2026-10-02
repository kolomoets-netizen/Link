#!/usr/bin/env python3
"""Build editable iStockLink Argon CRM deck as PPTX (16:9 B2B style)."""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

BLUE = RGBColor(0x1C, 0x50, 0xDE)
INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x66, 0x70, 0x85)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF5, 0xF6, 0xF8)
DARK = RGBColor(0x0B, 0x0B, 0x0B)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xEE, 0xF1, 0xF6)

W = Inches(13.333)
H = Inches(7.5)
MARGIN_X = Inches(0.7)
MARGIN_Y = Inches(0.45)


def set_run(run, size=18, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_multitext(slide, left, top, width, height, lines, size=16, color=MUTED, bold=False, spacing=1.1):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color)
        p.space_after = Pt(6)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def brand_bar(slide, dark=False):
    # blue mark
    mark = slide.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN_X, MARGIN_Y + Inches(0.05), Inches(0.16), Inches(0.16))
    fill_shape(mark, BLUE)
    add_textbox(
        slide,
        MARGIN_X + Inches(0.28),
        MARGIN_Y,
        Inches(3),
        Inches(0.3),
        "iStockLink",
        size=14,
        bold=True,
        color=WHITE if dark else INK,
    )


def footer(slide, meta, num, dark=False):
    c = RGBColor(0x99, 0x99, 0x99) if dark else MUTED
    add_textbox(slide, MARGIN_X, H - Inches(0.45), Inches(8), Inches(0.3), meta, size=11, color=c)
    add_textbox(
        slide,
        W - MARGIN_X - Inches(1.4),
        H - Inches(0.45),
        Inches(1.4),
        Inches(0.3),
        num,
        size=11,
        color=c,
        align=PP_ALIGN.RIGHT,
    )


def bg(slide, kind):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill_shape(shape, DARK if kind == "dark" else (PAPER if kind == "paper" else WHITE))
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shape, fill)
    try:
        shape.adjustments[0] = 0.1
    except Exception:
        pass
    return shape


def build(slides_data, out_path: Path):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    total = len(slides_data)

    for s in slides_data:
        slide = prs.slides.add_slide(blank)
        dark = s["kind"] == "dark"
        bg(slide, s["kind"])
        brand_bar(slide, dark=dark)
        title_color = WHITE if dark else INK
        muted = RGBColor(0xA8, 0xB0, 0xBD) if dark else MUTED
        num = f"{s['i']:02d} / {total:02d}"

        if s["kicker"]:
            add_textbox(
                slide,
                W - MARGIN_X - Inches(5.5),
                MARGIN_Y,
                Inches(5.5),
                Inches(0.3),
                s["kicker"].upper(),
                size=11,
                bold=True,
                color=BLUE if not dark else RGBColor(0x8E, 0xAF, 0xFF),
                align=PP_ALIGN.RIGHT,
            )

        # Title
        title = s["title"].replace("\n", " ")
        title_size = 40 if len(title) < 40 else (34 if len(title) < 70 else 28)
        add_textbox(
            slide,
            MARGIN_X,
            Inches(1.15),
            Inches(11.8),
            Inches(1.4),
            title,
            size=title_size,
            bold=True,
            color=title_color,
        )

        y = Inches(2.5)
        if s["lead"]:
            add_textbox(
                slide,
                MARGIN_X,
                Inches(2.45),
                Inches(11.5),
                Inches(0.9),
                s["lead"].replace("\n", " "),
                size=18,
                color=muted,
            )
            y = Inches(3.35)

        # TOC
        if s["toc"]:
            y = Inches(2.7)
            for row in s["toc"]:
                add_textbox(slide, MARGIN_X, y, Inches(0.7), Inches(0.35), row["n"], size=16, bold=True, color=BLUE)
                add_textbox(slide, MARGIN_X + Inches(0.8), y, Inches(6), Inches(0.35), row["t"], size=18, bold=True, color=INK)
                add_textbox(
                    slide,
                    MARGIN_X + Inches(7.2),
                    y,
                    Inches(4.5),
                    Inches(0.35),
                    row["d"],
                    size=14,
                    color=MUTED,
                    align=PP_ALIGN.RIGHT,
                )
                y += Inches(0.55)

        # Stats grid
        elif s["stats"]:
            n = len(s["stats"])
            gap = Inches(0.18)
            card_w = (W - 2 * MARGIN_X - gap * (n - 1)) / n
            card_h = Inches(3.2)
            top = Inches(3.3)
            for i, st in enumerate(s["stats"]):
                left = MARGIN_X + i * (card_w + gap)
                card(slide, left, top, card_w, card_h, WHITE if s["kind"] != "white" else SOFT)
                add_textbox(slide, left + Inches(0.2), top + Inches(0.2), card_w - Inches(0.3), Inches(0.3), f"{i+1:02d}", size=12, bold=True, color=BLUE)
                add_textbox(slide, left + Inches(0.2), top + Inches(0.55), card_w - Inches(0.3), Inches(0.6), st["big"], size=28, bold=True, color=INK)
                add_textbox(slide, left + Inches(0.2), top + Inches(1.3), card_w - Inches(0.3), Inches(0.7), st["h"], size=15, bold=True, color=INK)
                add_textbox(slide, left + Inches(0.2), top + Inches(2.1), card_w - Inches(0.3), Inches(0.9), st["p"], size=12, color=MUTED)

        # Items as cards / list
        elif s["items"]:
            items = s["items"]
            # special: left title already placed; put items in grid or two-col list
            if len(items) >= 6 and s["i"] in (5, 14, 15):
                cols = 3 if len(items) != 6 or s["i"] == 15 else 3
                if s["i"] == 15:
                    cols = 6
                rows = (len(items) + cols - 1) // cols
                gap = Inches(0.15)
                usable_w = W - 2 * MARGIN_X
                card_w = (usable_w - gap * (cols - 1)) / cols
                card_h = Inches(2.6) if cols < 6 else Inches(2.8)
                top = Inches(3.2)
                for idx, it in enumerate(items[: cols * rows]):
                    r, c = divmod(idx, cols) if cols != 6 else (0, idx)
                    if cols == 6:
                        left = MARGIN_X + idx * (card_w + gap)
                        top_i = top
                    else:
                        left = MARGIN_X + c * (card_w + gap)
                        top_i = top + r * (card_h + gap)
                    card(slide, left, top_i, card_w, card_h, WHITE if s["kind"] != "white" else SOFT)
                    add_textbox(slide, left + Inches(0.18), top_i + Inches(0.18), card_w - Inches(0.3), Inches(0.3), f"{idx+1:02d}", size=12, bold=True, color=BLUE)
                    add_textbox(slide, left + Inches(0.18), top_i + Inches(0.55), card_w - Inches(0.3), Inches(0.7), it["h"], size=14 if cols == 6 else 16, bold=True, color=INK)
                    add_textbox(slide, left + Inches(0.18), top_i + Inches(1.3), card_w - Inches(0.3), Inches(1.1), it["p"], size=11 if cols == 6 else 12, color=MUTED)
            elif s["i"] in (3, 23) and dark:
                # pain list on right-ish
                top = Inches(2.5)
                for idx, it in enumerate(items[:4]):
                    yy = top + idx * Inches(0.95)
                    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN_X + Inches(6.4), yy, Inches(0.42), Inches(0.42))
                    fill_shape(oval, RGBColor(0x1A, 0x2F, 0x6B))
                    add_textbox(slide, MARGIN_X + Inches(6.42), yy + Inches(0.05), Inches(0.4), Inches(0.35), f"{idx+1:02d}", size=11, bold=True, color=RGBColor(0x9C, 0xB6, 0xFF), align=PP_ALIGN.CENTER)
                    add_textbox(slide, MARGIN_X + Inches(7.0), yy, Inches(5.2), Inches(0.35), it["h"], size=16, bold=True, color=WHITE)
                    add_textbox(slide, MARGIN_X + Inches(7.0), yy + Inches(0.35), Inches(5.2), Inches(0.4), it["p"], size=13, color=muted)
                # shrink title column visually already ok
            elif s["i"] in (10, 20):
                top = Inches(2.6)
                for idx, it in enumerate(items[:4]):
                    yy = top + idx * Inches(0.95)
                    add_textbox(slide, MARGIN_X + Inches(6.2), yy, Inches(0.7), Inches(0.4), f"{idx+1:02d}", size=22, bold=True, color=BLUE)
                    add_textbox(slide, MARGIN_X + Inches(7.0), yy, Inches(5.2), Inches(0.35), it["h"], size=18, bold=True, color=INK)
                    add_textbox(slide, MARGIN_X + Inches(7.0), yy + Inches(0.35), Inches(5.2), Inches(0.4), it["p"], size=14, color=MUTED)
            else:
                # flow / generic cards in a row or stacked tasks
                if len(items) <= 4 and all(len(it["h"]) < 40 for it in items):
                    n = len(items)
                    gap = Inches(0.18)
                    card_w = (W - 2 * MARGIN_X - gap * (n - 1)) / n
                    top = Inches(3.4)
                    for i, it in enumerate(items):
                        left = MARGIN_X + i * (card_w + gap)
                        card(slide, left, top, card_w, Inches(2.8), WHITE if s["kind"] != "white" else SOFT)
                        add_textbox(slide, left + Inches(0.2), top + Inches(0.2), card_w - Inches(0.3), Inches(0.4), f"{i+1:02d}", size=20, bold=True, color=BLUE)
                        add_textbox(slide, left + Inches(0.2), top + Inches(0.8), card_w - Inches(0.3), Inches(0.6), it["h"], size=18, bold=True, color=INK)
                        add_textbox(slide, left + Inches(0.2), top + Inches(1.5), card_w - Inches(0.3), Inches(1.0), it["p"], size=13, color=MUTED)
                else:
                    top = Inches(3.2)
                    for idx, it in enumerate(items[:5]):
                        yy = top + idx * Inches(0.7)
                        card(slide, MARGIN_X, yy, W - 2 * MARGIN_X, Inches(0.62), WHITE if s["kind"] != "white" else SOFT)
                        add_textbox(slide, MARGIN_X + Inches(0.25), yy + Inches(0.12), Inches(8), Inches(0.4), it["h"], size=16, bold=True, color=INK)
                        add_textbox(
                            slide,
                            MARGIN_X + Inches(8.2),
                            yy + Inches(0.14),
                            Inches(3.8),
                            Inches(0.4),
                            it["p"],
                            size=12,
                            color=MUTED,
                            align=PP_ALIGN.RIGHT,
                        )

        # Cover / ending extras
        if s["i"] == 1:
            add_textbox(
                slide,
                MARGIN_X,
                Inches(0.95),
                Inches(4),
                Inches(0.3),
                "CRM ISTOCKLINK",
                size=12,
                bold=True,
                color=RGBColor(0x8E, 0xAF, 0xFF),
            )

        footer(slide, s.get("foot") or "iStockLink CRM", num, dark=dark)

    prs.save(str(out_path))
    print(f"saved {out_path} ({total} slides)")


if __name__ == "__main__":
    data = json.loads(Path("/tmp/slides.json").read_text(encoding="utf-8"))
    out = Path("/workspace/content/business/istocklink-argon-crm.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    build(data, out)
    Path("/workspace/docs/business").mkdir(parents=True, exist_ok=True)
    Path("/workspace/docs/business/istocklink-argon-crm.pptx").write_bytes(out.read_bytes())
