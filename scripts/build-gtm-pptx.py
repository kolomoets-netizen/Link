#!/usr/bin/env python3
"""Build editable GTM / business-plan PPTX for iStockLink CRM+SRM (not Argon sales deck)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BLUE = RGBColor(0x1C, 0x50, 0xDE)
INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x66, 0x70, 0x85)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF5, 0xF6, 0xF8)
DARK = RGBColor(0x0B, 0x0B, 0x0B)
SOFT = RGBColor(0xEE, 0xF1, 0xF6)
BLUE_SOFT = RGBColor(0xEA, 0xF0, 0xFF)

W = Inches(13.333)
H = Inches(7.5)
MX = Inches(0.7)
MY = Inches(0.45)


def set_run(run, size=18, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def bg(slide, kind="white"):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(shape, DARK if kind == "dark" else (PAPER if kind == "paper" else WHITE))
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def multilines(slide, left, top, width, height, lines, size=14, color=MUTED, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color)
        p.space_after = Pt(4)
    return box


def card(slide, left, top, width, height, color=SOFT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill(shape, color)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def brand(slide, dark=False):
    mark = slide.shapes.add_shape(MSO_SHAPE.OVAL, MX, MY + Inches(0.04), Inches(0.16), Inches(0.16))
    fill(mark, BLUE)
    textbox(slide, MX + Inches(0.28), MY, Inches(3), Inches(0.28), "iStockLink", 14, True, WHITE if dark else INK)


def kicker(slide, text, dark=False):
    textbox(
        slide,
        W - MX - Inches(6),
        MY,
        Inches(6),
        Inches(0.28),
        text.upper(),
        11,
        True,
        RGBColor(0x8E, 0xAF, 0xFF) if dark else BLUE,
        PP_ALIGN.RIGHT,
    )


def footer(slide, meta, num, total, dark=False):
    c = RGBColor(0x99, 0x99, 0x99) if dark else MUTED
    textbox(slide, MX, H - Inches(0.42), Inches(8), Inches(0.28), meta, 11, False, c)
    textbox(slide, W - MX - Inches(1.5), H - Inches(0.42), Inches(1.5), Inches(0.28), f"{num:02d} / {total:02d}", 11, False, c, PP_ALIGN.RIGHT)


def title(slide, text, dark=False, size=34, top=Inches(1.1)):
    textbox(slide, MX, top, Inches(12), Inches(1.1), text, size, True, WHITE if dark else INK)


def lead(slide, text, dark=False, top=Inches(2.2)):
    textbox(slide, MX, top, Inches(11.5), Inches(0.8), text, 16, False, RGBColor(0xA8, 0xB0, 0xBD) if dark else MUTED)


def add_table(slide, left, top, width, height, rows, col_widths=None, header=True):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, height)
    table = table_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    is_header = header and r == 0
                    set_run(run, size=10 if not is_header else 10, bold=is_header or c == 0, color=WHITE if is_header else INK)
            # fill
            fill_color = DARK if (header and r == 0) else (SOFT if r % 2 == 0 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill_color
    return table_shape


def build(out: Path):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    slides_meta = []

    def new(kind="white"):
        s = prs.slides.add_slide(blank)
        bg(s, kind)
        return s

    # 01 Cover
    s = new("dark")
    brand(s, True)
    kicker(s, "Go-to-market · Confidential", True)
    textbox(s, MX, Inches(1.8), Inches(10), Inches(0.35), "БИЗНЕС-ПЛАН ПРОДВИЖЕНИЯ", 12, True, RGBColor(0x8E, 0xAF, 0xFF))
    textbox(s, MX, Inches(2.3), Inches(12), Inches(1.8), "Анализ рынка и план\nпродвижения CRM + SRM", 40, True, WHITE)
    lead(s, "Деньги в каналы → SQL → оплаты. Калибровка на Яндекс Директе, сентябрь 2026.", True, Inches(4.4))
    card(s, MX, Inches(5.3), Inches(7.5), Inches(1.2), RGBColor(0x17, 0x17, 0x17))
    multilines(
        s,
        MX + Inches(0.25),
        Inches(5.45),
        Inches(7),
        Inches(1.0),
        ["Продукт: iStockLink CRM + SRM", "Факт: CPL 3 447 ₽ · SQL 5 515 ₽ (сент. 2026)", "Версия 1.2 · октябрь 2026"],
        12,
        RGBColor(0xC5, 0xCB, 0xD3),
    )
    slides_meta.append(("Сборка GTM", "dark"))

    # 02 Agenda
    s = new("white")
    brand(s)
    kicker(s, "Содержание")
    title(s, "О чём этот файл")
    items = [
        ("01", "Резюме", "Рынок, окно, фокус на 12 месяцев"),
        ("02", "Анализ рынка", "CRM, SRM, ICP, конкуренты"),
        ("03", "Деньги → результат", "Сценарии бюджета A / B / C"),
        ("04", "С Битрикс24", "Функции, TCO, когда что выбирать"),
        ("05", "Что нужно от вас", "Выгрузки рекламы и касаний"),
    ]
    y = Inches(2.5)
    for n, t, d in items:
        textbox(s, MX, y, Inches(0.8), Inches(0.4), n, 18, True, BLUE)
        textbox(s, MX + Inches(1.0), y, Inches(4), Inches(0.4), t, 20, True, INK)
        textbox(s, MX + Inches(5.5), y, Inches(6.5), Inches(0.4), d, 16, False, MUTED, PP_ALIGN.RIGHT)
        y += Inches(0.7)
    slides_meta.append(("Содержание", "white"))

    # 03 Exec
    s = new("paper")
    brand(s)
    kicker(s, "Резюме")
    title(s, "Ниша между CRM-комбайном и тяжёлым SRM")
    lead(s, "Один контур: сделка → запрос КП → поставка. Без проекта внедрения.")
    kpis = [
        ("~44 млрд ₽", "Рынок CRM РФ, 2025\nTAdviser, +25% г/г"),
        ("~3,5 млрд ₽", "Рынок SRM РФ, 2025\nKept / TAdviser"),
        ("1 день", "Старт iStockLink\nvs недели у Битрикс + внедрение"),
    ]
    for i, (n, l) in enumerate(kpis):
        left = MX + i * Inches(4.05)
        card(s, left, Inches(3.3), Inches(3.85), Inches(2.6), WHITE)
        textbox(s, left + Inches(0.25), Inches(3.55), Inches(3.4), Inches(0.7), n, 26, True, BLUE)
        multilines(s, left + Inches(0.25), Inches(4.4), Inches(3.4), Inches(1.2), l.split("\n"), 13, MUTED)
    slides_meta.append(("Резюме", "paper"))

    # 04 Conclusions
    s = new("white")
    brand(s)
    kicker(s, "Резюме")
    title(s, "Ключевые выводы")
    points = [
        "Спрос на отечественные CRM и закупки растёт.",
        "Средний бизнес платит за Битрикс дважды: лицензия + внедрение 150–600 тыс. ₽.",
        "SRM растёт быстрее CRM, но тяжёл для mid-market.",
        "Окно iStockLink: 15–150 сотрудников, продажи + регулярные закупки/торги.",
    ]
    y = Inches(2.5)
    for i, p in enumerate(points):
        card(s, MX, y, W - 2 * MX, Inches(0.85), SOFT)
        textbox(s, MX + Inches(0.25), y + Inches(0.2), Inches(0.6), Inches(0.4), f"{i+1:02d}", 16, True, BLUE)
        textbox(s, MX + Inches(1.0), y + Inches(0.22), Inches(10.8), Inches(0.5), p, 16, False, INK)
        y += Inches(1.0)
    slides_meta.append(("Выводы", "white"))

    # 05 Market CRM
    s = new("paper")
    brand(s)
    kicker(s, "Рынок")
    title(s, "Рынок CRM в России")
    add_table(
        s,
        MX,
        Inches(2.4),
        W - 2 * MX,
        Inches(3.8),
        [
            ["Показатель", "Оценка", "Источник"],
            ["Объём CRM, 2024", "~35,3 млрд ₽", "TAdviser"],
            ["Объём CRM, 2025", "~44,1 млрд ₽ (+25%)", "TAdviser"],
            ["Комплексные CRM, 2025", "~20,5 млрд ₽", "Kept"],
            ["Прогноз к 2032", "~46,1 млрд ₽", "Kept, CAGR ~12%"],
            ["Лидеры восприятия", "Битрикс24, amoCRM, 1С…", "Универсальные платформы"],
        ],
        col_widths=[Inches(4.2), Inches(4.0), Inches(3.8)],
    )
    slides_meta.append(("Рынок CRM", "paper"))

    # 06 Market SRM
    s = new("white")
    brand(s)
    kicker(s, "Рынок")
    title(s, "Рынок SRM и автоматизации закупок")
    add_table(
        s,
        MX,
        Inches(2.4),
        W - 2 * MX,
        Inches(3.6),
        [
            ["Показатель", "Оценка", "Источник"],
            ["SRM РФ, 2025", "~3,5 млрд ₽", "Kept / TAdviser"],
            ["CAGR SRM 2020–2025", "~30,6%", "Kept"],
            ["Автоматизация закупок, 2024", "~11,5 млрд ₽", "Data Insight"],
            ["Прогноз к 2030", "~18,7 млрд ₽", "Data Insight"],
            ["Penetration спецПО закупок", "~29% mid/large", "Росстат / Kept"],
        ],
        col_widths=[Inches(4.5), Inches(3.8), Inches(3.7)],
    )
    slides_meta.append(("Рынок SRM", "white"))

    # 07 ICP
    s = new("paper")
    brand(s)
    kicker(s, "ICP")
    title(s, "Кому продаём")
    card(s, MX, Inches(2.4), Inches(5.8), Inches(4.0), WHITE)
    textbox(s, MX + Inches(0.3), Inches(2.6), Inches(5.2), Inches(0.4), "Primary ICP", 16, True, BLUE)
    multilines(
        s,
        MX + Inches(0.3),
        Inches(3.2),
        Inches(5.2),
        Inches(3.0),
        [
            "• B2B / B2G, 15–150 сотрудников",
            "• Отдел продаж + регулярные закупки",
            "• Сейчас: Excel / «сырой» Битрикс24",
            "• Боли: потеря сделок, ручной сбор КП",
        ],
        15,
        INK,
    )
    card(s, MX + Inches(6.1), Inches(2.4), Inches(5.8), Inches(4.0), WHITE)
    textbox(s, MX + Inches(6.4), Inches(2.6), Inches(5.2), Inches(0.4), "Secondary ICP", 16, True, BLUE)
    multilines(
        s,
        MX + Inches(6.4),
        Inches(3.2),
        Inches(5.2),
        Inches(3.0),
        [
            "• Торговля / производство под клиента",
            "• Тендерные команды in-house / outsource",
            "• Переход с Битрикс без бюджета на интегратора",
            "• Нужны Контур / Seldon в том же контуре",
        ],
        15,
        INK,
    )
    slides_meta.append(("ICP", "paper"))

    # 08 Competition
    s = new("white")
    brand(s)
    kicker(s, "Конкуренты")
    title(s, "Конкурентная карта")
    add_table(
        s,
        MX,
        Inches(2.3),
        W - 2 * MX,
        Inches(4.2),
        [
            ["Кластер", "Примеры", "Слабость vs iStockLink"],
            ["Универсальные CRM", "Битрикс24, amoCRM", "SRM/КП — доработка; долгое внедрение"],
            ["Тяжёлые SRM / S2P", "Norbit, ELMA365…", "Дорого, слабо связано с продажами SMB"],
            ["Агрегаторы тендеров", "Контур, Seldon", "Нет CRM/SRM-операционки"],
            ["Excel + почта", "Статус-кво", "Нет контроля и масштаба"],
        ],
        col_widths=[Inches(3.5), Inches(3.8), Inches(4.7)],
    )
    slides_meta.append(("Конкуренты", "white"))

    # 09 Positioning
    s = new("dark")
    brand(s, True)
    kicker(s, "Позиционирование", True)
    title(s, "CRM + SRM в одном окне", True, 38, Inches(2.2))
    lead(s, "Не «ещё одна CRM». Снижение TCO и ускорение цикла «клиент → поставка».", True, Inches(3.6))
    textbox(s, MX, Inches(4.8), Inches(12), Inches(0.5), "Против: Excel · пустой Битрикс24 · разрозненные агрегаторы", 16, False, RGBColor(0x8E, 0xAF, 0xFF))
    slides_meta.append(("Позиционирование", "dark"))

    # 10 Goals
    s = new("paper")
    brand(s)
    kicker(s, "Цели 12 мес.")
    title(s, "Цели продвижения")
    add_table(
        s,
        MX,
        Inches(2.4),
        W - 2 * MX,
        Inches(4.0),
        [
            ["Цель", "Метрика", "Ориентир"],
            ["Квалифицированные демо", "SQL / мес к M12", "40–60"],
            ["Конверсия демо → оплата", "Win rate", "20–30%"],
            ["Новые платящие", "Paid logos / год", "80–120"],
            ["Переходы с Битрикс24", "% новых клиентов", "≥25%"],
            ["Удержание", "Logo retention 12 мес.", "≥85%"],
        ],
        col_widths=[Inches(4.5), Inches(4.2), Inches(3.3)],
    )
    slides_meta.append(("Цели", "paper"))

    # 11 Fact Sept
    s = new("dark")
    brand(s, True)
    kicker(s, "Факт · сентябрь 2026", True)
    title(s, "Яндекс Директ: что уже работает", True, 34, Inches(1.5))
    facts = [
        ("27 576 ₽", "расход за месяц"),
        ("8 лидов", "CPL 3 447 ₽"),
        ("5 SQL", "5 515 ₽ / SQL"),
        ("62%", "Lead → SQL"),
    ]
    for i, (n, l) in enumerate(facts):
        left = MX + (i % 4) * Inches(3.05)
        card(s, left, Inches(3.3), Inches(2.9), Inches(2.4), RGBColor(0x17, 0x17, 0x17))
        textbox(s, left + Inches(0.2), Inches(3.6), Inches(2.5), Inches(0.7), n, 22, True, BLUE)
        textbox(s, left + Inches(0.2), Inches(4.5), Inches(2.5), Inches(0.8), l, 14, False, RGBColor(0xA8, 0xB0, 0xBD))
    slides_meta.append(("Факт Директа", "dark"))

    # 12 Assumptions
    s = new("white")
    brand(s)
    kicker(s, "Модель")
    title(s, "Допущения после калибровки")
    add_table(
        s,
        MX,
        Inches(2.4),
        W - 2 * MX,
        Inches(4.0),
        [
            ["Параметр", "Значение", "Комментарий"],
            ["ACV команды", "120 000 ₽ / год", "Уточнить по факту сделок"],
            ["Cost / SQL сейчас", "~5 500 ₽", "Факт сентября"],
            ["Cost / SQL при масштабе", "6–10 тыс. ₽", "Рост аукциона"],
            ["Lead → SQL", "62%", "5 из 8"],
            ["SQL → оплата", "15% / 22% / 28%", "Факта закрытий пока нет"],
        ],
        col_widths=[Inches(3.8), Inches(3.5), Inches(4.7)],
    )
    slides_meta.append(("Допущения", "white"))

    # 13 Scenarios
    s = new("paper")
    brand(s)
    kicker(s, "Сценарии")
    title(s, "Если столько денег — такой результат")
    scenarios = [
        ("A · Разогрев", "150 тыс ₽/мес", "SQL ~300 / год", "~66 оплат*", "~7,9 млн ₽", "Рекомендуемый следующий шаг"),
        ("B · База", "300 тыс ₽/мес", "SQL ~480 / год", "~106 оплат*", "~12,7 млн ₽", "После фиксации win rate"),
        ("C · Рост", "600 тыс ₽/мес", "SQL ~720 / год", "~158 оплат*", "~19 млн ₽", "Нужны сайт + 2+ продавца"),
    ]
    for i, (name, m, sql, deals, rev, chance) in enumerate(scenarios):
        left = MX + i * Inches(4.05)
        card(s, left, Inches(2.35), Inches(3.85), Inches(4.3), WHITE)
        textbox(s, left + Inches(0.25), Inches(2.55), Inches(3.4), Inches(0.4), name, 18, True, BLUE)
        textbox(s, left + Inches(0.25), Inches(3.15), Inches(3.4), Inches(0.45), m, 18, True, INK)
        multilines(
            s,
            left + Inches(0.25),
            Inches(3.8),
            Inches(3.4),
            Inches(2.5),
            [sql, deals + " при win 22%", "Выручка: " + rev, chance],
            13,
            MUTED,
        )
    slides_meta.append(("Сценарии A/B/C", "paper"))

    # 14 Mix
    s = new("white")
    brand(s)
    kicker(s, "Каналы")
    title(s, "Куда класть деньги внутри бюджета")
    add_table(
        s,
        MX,
        Inches(2.3),
        W - 2 * MX,
        Inches(4.3),
        [
            ["Канал", "A Тест", "B База", "C Рост", "Зачем"],
            ["Яндекс Директ — поиск", "55%", "45%", "40%", "Заявки на демо"],
            ["Ретаргет", "15%", "15%", "15%", "Догон сайта"],
            ["Контент / кейсы / PDF", "15%", "15%", "12%", "Доверие + sales kit"],
            ["Вебинары / посевы", "5%", "10%", "12%", "Квалифицированный inbound"],
            ["Партнёры", "5%", "10%", "15%", "Агентства, 1С"],
            ["Тесты каналов", "5%", "5%", "6%", "Telegram / ABM"],
        ],
        col_widths=[Inches(3.4), Inches(1.5), Inches(1.5), Inches(1.5), Inches(3.9)],
    )
    slides_meta.append(("Микс каналов", "white"))

    # 15 Unit economics
    s = new("paper")
    brand(s)
    kicker(s, "Unit-экономика")
    title(s, "Когда можно масштабировать")
    rules = [
        ("CAC", "≤ 35–45% ACV year-1", "Иначе резать дорогие связки"),
        ("Payback", "≤ 6–9 месяцев", "Иначе поднимать ACV / сокращать цикл"),
        ("CPL стабильность", "4–8 недель", "Не лить search без ≥30 конверсий"),
        ("Win rate демо", "≥ 20%", "Сначала чинить sales, не бюджет"),
    ]
    for i, (h, v, note) in enumerate(rules):
        left = MX + (i % 2) * Inches(6.15)
        top = Inches(2.4) + (i // 2) * Inches(2.0)
        card(s, left, top, Inches(5.9), Inches(1.8), WHITE)
        textbox(s, left + Inches(0.3), top + Inches(0.25), Inches(5.3), Inches(0.35), h, 14, True, BLUE)
        textbox(s, left + Inches(0.3), top + Inches(0.7), Inches(5.3), Inches(0.4), v, 18, True, INK)
        textbox(s, left + Inches(0.3), top + Inches(1.2), Inches(5.3), Inches(0.35), note, 13, False, MUTED)
    slides_meta.append(("Unit-экономика", "paper"))

    # 16 Bitrix compare
    s = new("white")
    brand(s)
    kicker(s, "Битрикс24")
    title(s, "Сравнение с Битрикс24")
    add_table(
        s,
        MX,
        Inches(2.2),
        W - 2 * MX,
        Inches(4.5),
        [
            ["Критерий", "Битрикс24 + внедрение", "iStockLink CRM+SRM"],
            ["Время до процесса", "Недели–месяцы", "Около 1 дня"],
            ["SRM / поставщики", "Доработка", "В том же окне"],
            ["КП и сравнение", "Почта / кастом", "Сводная таблица"],
            ["Сделка → закупка", "Проектируется", "Единый маршрут"],
            ["Агрегаторы", "Через интегратора", "Контур / Seldon из коробки"],
            ["Модель затрат", "Подписка + 150–600+ тыс.", "Подписка, без проекта"],
        ],
        col_widths=[Inches(3.5), Inches(4.5), Inches(4.0)],
    )
    slides_meta.append(("Сравнение Битрикс", "white"))

    # 17 When choose
    s = new("paper")
    brand(s)
    kicker(s, "Битрикс24")
    title(s, "Когда что выбирать")
    card(s, MX, Inches(2.4), Inches(5.8), Inches(4.0), WHITE)
    textbox(s, MX + Inches(0.3), Inches(2.65), Inches(5.2), Inches(0.4), "Битрикс24", 18, True, MUTED)
    multilines(
        s,
        MX + Inches(0.3),
        Inches(3.3),
        Inches(5.2),
        Inches(2.8),
        [
            "• Нужен «комбайн»: CRM + портал + задачи",
            "• Есть бюджет и время на интегратора",
            "• Закупки — не ключевой процесс",
        ],
        15,
        INK,
    )
    card(s, MX + Inches(6.1), Inches(2.4), Inches(5.8), Inches(4.0), BLUE_SOFT)
    textbox(s, MX + Inches(6.4), Inches(2.65), Inches(5.2), Inches(0.4), "iStockLink", 18, True, BLUE)
    multilines(
        s,
        MX + Inches(6.4),
        Inches(3.3),
        Inches(5.2),
        Inches(2.8),
        [
            "• Продажи и поставщики в одном цикле",
            "• Нужен быстрый старт без проекта",
            "• Важны КП, SRM, тендеры в том же окне",
        ],
        15,
        INK,
    )
    slides_meta.append(("Когда выбирать", "paper"))

    # 18 Data ask
    s = new("dark")
    brand(s, True)
    kicker(s, "Данные", True)
    title(s, "Директ учтён. Не хватает закрытий", True, 30, Inches(1.5))
    blocks = [
        ("1", "CRM: SQL → оплата", "По 5 SQL сентября — статус, тариф, won/lost"),
        ("2", "Касания 3–6 мес.", "Источник / UTM, этап, сумма — для win rate"),
        ("3", "Средний чек", "Факт ACV по оплаченным аккаунтам"),
    ]
    for i, (n, h, p) in enumerate(blocks):
        top = Inches(3.0) + i * Inches(1.15)
        textbox(s, MX, top, Inches(0.6), Inches(0.4), n, 22, True, BLUE)
        textbox(s, MX + Inches(0.8), top, Inches(11), Inches(0.4), h, 20, True, WHITE)
        textbox(s, MX + Inches(0.8), top + Inches(0.4), Inches(11), Inches(0.4), p, 14, False, RGBColor(0xA8, 0xB0, 0xBD))
    slides_meta.append(("Нужные данные", "dark"))

    # 19 Close
    s = new("dark")
    brand(s, True)
    kicker(s, "Дальше", True)
    title(s, "Следующий шаг: сценарий A", True, 34, Inches(2.2))
    lead(s, "150 тыс ₽/мес · ~300 SQL/год · ~66 оплат при win 22% · ~7,9 млн ₽ (модель).", True, Inches(3.7))
    textbox(s, MX, Inches(4.8), Inches(12), Inches(0.5), "Сейчас ~28 тыс ₽/мес. Сначала ×5 и замер win rate, потом B/C.", 16, False, RGBColor(0x8E, 0xAF, 0xFF))
    textbox(s, MX, Inches(5.6), Inches(12), Inches(0.4), "istock.link", 20, True, WHITE)
    slides_meta.append(("Рекомендация", "dark"))

    total = len(prs.slides)
    # add footers
    for i, slide in enumerate(prs.slides):
        meta, kind = slides_meta[i]
        dark = kind == "dark"
        # skip if already has lots - just add footer
        footer(slide, meta if i else "iStockLink · GTM CRM+SRM", i + 1, total, dark)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"saved {out} ({total} slides)")


if __name__ == "__main__":
    out = Path("/workspace/content/business/istocklink-crm-srm-gtm.pptx")
    build(out)
    docs = Path("/workspace/docs/business/istocklink-crm-srm-gtm.pptx")
    docs.write_bytes(out.read_bytes())
    Path("/opt/cursor/artifacts/istocklink_crm_srm_gtm.pptx").write_bytes(out.read_bytes())
